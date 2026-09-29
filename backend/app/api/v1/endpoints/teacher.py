from fastapi import APIRouter, Depends, HTTPException, status
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_
from typing import List, Optional

from app.db.session import get_db
from app.core.dependencies import get_current_user
from app.models.user import User, UserRole
from app.models.academic import (
    Module, TimeSlot, ExamSchedule, Section, Level, Semester,
    teacher_modules, student_enrollments
)
from app.schemas.academic import TimeSlotRead, ExamScheduleRead, ModuleRead
from app.schemas.user import UserRead

router = APIRouter(prefix="/teacher", tags=["👨‍🏫 Espace Enseignant"])


def _require_teacher(current_user: User = Depends(get_current_user)) -> User:
    if current_user.role not in (UserRole.TEACHER, UserRole.ADMIN):
        raise HTTPException(status_code=403, detail="Accès réservé aux enseignants")
    return current_user


def _get_teacher_module_ids(db: Session, teacher_id: int) -> List[int]:
    """Helper to get all module IDs for a teacher, either directly assigned or in their timetable."""
    # 1. Modules directly assigned
    assigned_rows = db.query(teacher_modules.c.module_id).filter(teacher_modules.c.teacher_id == teacher_id).all()
    # 2. Modules from timetable
    timetable_rows = db.query(TimeSlot.module_id).filter(TimeSlot.teacher_id == teacher_id).all()
    
    # Combine and deduplicate, ignoring None
    all_ids = set(
        [r[0] for r in assigned_rows if r[0] is not None] + 
        [r[0] for r in timetable_rows if r[0] is not None]
    )
    return list(all_ids)


# ── Mes Modules ────────────────────────────────────────────────────────────────
@router.get("/me/modules")
def get_my_modules(
    db: Session = Depends(get_db),
    current_user: User = Depends(_require_teacher),
):
    """Retourne les modules enseignés par le professeur connecté, avec contexte académique."""
    mod_ids = _get_teacher_module_ids(db, current_user.id)
    if not mod_ids:
        return []

    modules = (
        db.query(Module)
        .filter(Module.id.in_(mod_ids))
        .options(joinedload(Module.semester).joinedload(Semester.level).joinedload(Level.specialty))
        .all()
    )
    
    result = []
    for m in modules:
        sem = m.semester
        lvl = sem.level if sem else None
        spec = lvl.specialty if lvl else None
        result.append({
            "id": m.id,
            "name": m.name,
            "semester": {"id": sem.id, "name": sem.name} if sem else None,
            "level": {"id": lvl.id, "name": lvl.name} if lvl else None,
            "specialty": {"id": spec.id, "name": spec.name} if spec else None,
        })
    return result


# ── Mon Emploi du Temps ────────────────────────────────────────────────────────
@router.get("/me/timetable", response_model=List[TimeSlotRead])
def get_my_timetable(
    db: Session = Depends(get_db),
    current_user: User = Depends(_require_teacher),
):
    """Retourne toutes les séances (cours, TD, TP) du professeur connecté."""
    return (
        db.query(TimeSlot)
        .filter(TimeSlot.teacher_id == current_user.id)
        .options(
            joinedload(TimeSlot.module),
            joinedload(TimeSlot.section),
            joinedload(TimeSlot.group),
        )
        .all()
    )


# ── Mes Étudiants ──────────────────────────────────────────────────────────────
@router.get("/me/students")
def get_my_students(
    db: Session = Depends(get_db),
    current_user: User = Depends(_require_teacher),
):
    """
    Retourne les étudiants inscrits dans les niveaux correspondants
    aux modules enseignés par le professeur, regroupés par module.
    """
    mod_ids = _get_teacher_module_ids(db, current_user.id)
    if not mod_ids:
        return []

    modules = (
        db.query(Module)
        .filter(Module.id.in_(mod_ids))
        .options(joinedload(Module.semester).joinedload(Semester.level))
        .all()
    )

    seen_levels = {}
    result = []

    for m in modules:
        sem = m.semester
        lvl = sem.level if sem else None
        if not lvl:
            continue

        if lvl.id not in seen_levels:
            students = (
                db.query(User)
                .join(student_enrollments, User.id == student_enrollments.c.student_id)
                .filter(student_enrollments.c.level_id == lvl.id)
                .all()
            )
            seen_levels[lvl.id] = [
                {
                    "id": s.id,
                    "first_name": s.first_name,
                    "last_name": s.last_name,
                    "email": s.email,
                } for s in students
            ]

        result.append({
            "module_id": m.id,
            "module_name": m.name,
            "level_id": lvl.id,
            "level_name": lvl.name,
            "students": seen_levels[lvl.id],
        })

    return result


# ── Mes Examens ────────────────────────────────────────────────────────────────
@router.get("/me/exams", response_model=List[ExamScheduleRead])
def get_my_exams(
    db: Session = Depends(get_db),
    current_user: User = Depends(_require_teacher),
):
    """Retourne les examens planifiés pour les modules du professeur."""
    mod_ids = _get_teacher_module_ids(db, current_user.id)
    if not mod_ids:
        return []

    return (
        db.query(ExamSchedule)
        .filter(ExamSchedule.module_id.in_(mod_ids))
        .options(
            joinedload(ExamSchedule.module),
            joinedload(ExamSchedule.level),
            joinedload(ExamSchedule.section),
            joinedload(ExamSchedule.uploaded_by),
        )
        .order_by(ExamSchedule.exam_date, ExamSchedule.start_time)
        .all()
    )


# ── Upload Sujet d'Examen ──────────────────────────────────────────────────────
from fastapi import UploadFile, File
import shutil
import os
import uuid

@router.post("/me/exams/{exam_id}/upload", response_model=ExamScheduleRead)
async def upload_exam_file(
    exam_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(_require_teacher),
):
    """Le professeur uploade le fichier sujet d'examen pour un examen programmé."""
    # Vérifier que l'examen existe
    exam = db.query(ExamSchedule).filter(ExamSchedule.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Examen introuvable.")

    # Vérifier que cet examen appartient à un module du professeur
    mod_ids = _get_teacher_module_ids(db, current_user.id)
    if exam.module_id not in mod_ids:
        raise HTTPException(status_code=403, detail="Cet examen n'appartient pas à vos modules.")

    # Valider l'extension du fichier
    allowed_extensions = {'.pdf', '.doc', '.docx', '.odt', '.zip', '.rar'}
    file_ext = os.path.splitext(file.filename)[1].lower()
    if file_ext not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail=f"Type de fichier non autorisé. Formats acceptés: PDF, DOC, DOCX, ODT, ZIP, RAR"
        )

    # Sauvegarder le fichier
    os.makedirs("uploads/exams", exist_ok=True)
    unique_filename = f"exam_{exam_id}_{uuid.uuid4().hex[:8]}{file_ext}"
    file_path = os.path.join("uploads", "exams", unique_filename)

    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    # Mettre à jour l'enregistrement en base
    exam.exam_file_url = f"/uploads/exams/{unique_filename}"
    exam.exam_file_name = file.filename
    exam.uploaded_by_id = current_user.id
    db.commit()
    db.refresh(exam)

    return exam


# ── Statistiques rapides ──────────────────────────────────────────────────────
@router.get("/me/stats")
def get_my_stats(
    db: Session = Depends(get_db),
    current_user: User = Depends(_require_teacher),
):
    """Statistiques pour le dashboard enseignant."""
    mod_ids = _get_teacher_module_ids(db, current_user.id)
    module_count = len(mod_ids)

    session_count = (
        db.query(TimeSlot)
        .filter(TimeSlot.teacher_id == current_user.id)
        .count()
    )

    student_count = 0
    exam_count = 0
    if mod_ids:
        exam_count = db.query(ExamSchedule).filter(ExamSchedule.module_id.in_(mod_ids)).count()

    return {
        "modules": module_count,
        "sessions": session_count,
        "students": student_count,
        "exams": exam_count,
    }


# ── Gestion des Notes ──────────────────────────────────────────────────────────
from fastapi import UploadFile, File, Form
from io import BytesIO
import openpyxl
from app.models.academic import Grade, GradeType
from app.schemas.academic import GradeRead, GradeBulkUploadResponse


@router.post("/me/grades/upload-excel", response_model=GradeBulkUploadResponse)
async def upload_grades_excel(
    module_id: int = Form(...),
    grade_type: str = Form(...),  # "exam" | "td" | "tp"
    academic_year: str = Form(default="2024-2025"),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(_require_teacher),
@router.get("/me/modules/{module_id}/students-template")
def get_module_students_template(
    module_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(_require_teacher),
):
    """
    Retourne la liste des étudiants inscrits au niveau de ce module
    pour pré-remplir le modèle Excel du professeur.
    """
    mod_ids = _get_teacher_module_ids(db, current_user.id)
    if module_id not in mod_ids:
        raise HTTPException(status_code=403, detail="Ce module ne vous est pas affecté.")

    module = db.query(Module).filter(Module.id == module_id).options(
        joinedload(Module.semester).joinedload(Semester.level)
    ).first()
    if not module or not module.semester or not module.semester.level:
        return []

    level_id = module.semester.level.id
    students = (
        db.query(User)
        .join(student_enrollments, User.id == student_enrollments.c.student_id)
        .filter(student_enrollments.c.level_id == level_id)
        .order_by(User.last_name, User.first_name)
        .all()
    )

    return [
        {
            "id": s.id,
            "last_name": s.last_name,
            "first_name": s.first_name,
            "email": s.email,
        }
        for s in students
    ]


@router.post("/me/grades/upload-excel", response_model=GradeBulkUploadResponse)
async def upload_grades_excel(
    module_id: int = Form(...),
    grade_type: str = Form(...),  # "exam" | "td" | "tp"
    academic_year: str = Form(default="2024-2025"),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(_require_teacher),
):
    """
    Upload des notes depuis un fichier Excel avec détection intelligente :
    Supporte :
    - ID / Numéro étudiant (ex: 12)
    - Nom + Prénom (ex: Benali Ahmed)
    - Email (ex: ahmed@univ.dz)
    - Colonne Note (0-20)
    """
    # Valider le type de note
    try:
        gtype = GradeType(grade_type.lower())
    except ValueError:
        raise HTTPException(status_code=400, detail="grade_type doit être 'exam', 'td' ou 'tp'")

    # Vérifier que le module appartient au professeur
    mod_ids = _get_teacher_module_ids(db, current_user.id)
    if module_id not in mod_ids:
        raise HTTPException(status_code=403, detail="Ce module ne vous est pas affecté.")

    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Veuillez fournir un fichier Excel (.xlsx ou .xls)")

    try:
        contents = await file.read()
        wb = openpyxl.load_workbook(filename=BytesIO(contents), data_only=True)
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Impossible de lire le fichier Excel : {str(e)}")

    # Pré-charger tous les étudiants de l'université pour recherche ultra rapide
    all_students = db.query(User).filter(
        User.university_id == current_user.university_id,
        User.role == UserRole.STUDENT
    ).all()

    students_by_id = {s.id: s for s in all_students}
    students_by_email = {s.email.lower().strip(): s for s in all_students if s.email}
    students_by_name = {f"{s.last_name.lower().strip()} {s.first_name.lower().strip()}": s for s in all_students}
    students_by_name_rev = {f"{s.first_name.lower().strip()} {s.last_name.lower().strip()}": s for s in all_students}

    # Analyser les en-têtes (ligne 1)
    header_row = [str(cell).strip().lower() if cell is not None else "" for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    
    col_id_idx = None
    col_nom_idx = None
    col_prenom_idx = None
    col_email_idx = None
    col_note_idx = None

    for idx, h in enumerate(header_row):
        if any(kw in h for kw in ['id', 'matricule', 'num', 'numero', 'code']):
            col_id_idx = idx
        elif any(kw in h for kw in ['prenom', 'first_name', 'firstname']):
            col_prenom_idx = idx
        elif any(kw in h for kw in ['nom', 'last_name', 'lastname']):
            col_nom_idx = idx
        elif 'email' in h or 'mail' in h:
            col_email_idx = idx
        elif any(kw in h for kw in ['note', 'score', 'examen', 'eval', 'mark']):
            col_note_idx = idx

    imported = 0
    updated = 0
    errors = 0
    error_details = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        # Trouver la note et l'étudiant
        student = None
        score_val = None

        # 1. Si les colonnes d'en-tête ont été identifiées
        if col_note_idx is not None and col_note_idx < len(row):
            score_val = row[col_note_idx]

        # Recherche de l'étudiant par ID
        if col_id_idx is not None and col_id_idx < len(row) and row[col_id_idx] is not None:
            try:
                sid = int(float(str(row[col_id_idx]).strip()))
                student = students_by_id.get(sid)
            except (ValueError, TypeError):
                pass

        # Recherche par email si pas trouvé
        if not student and col_email_idx is not None and col_email_idx < len(row) and row[col_email_idx]:
            em = str(row[col_email_idx]).strip().lower()
            student = students_by_email.get(em)

        # Recherche par Nom + Prénom si pas trouvé
        if not student and col_nom_idx is not None and col_prenom_idx is not None:
            nom = str(row[col_nom_idx] or "").strip().lower()
            prenom = str(row[col_prenom_idx] or "").strip().lower()
            student = students_by_name.get(f"{nom} {prenom}") or students_by_name_rev.get(f"{prenom} {nom}")

        # 2. Si pas d'en-tête formel, détection automatique selon le nombre de colonnes
        if not student:
            # Essai colonne 0 comme ID ou email ou Nom
            c0 = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
            c1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            
            # Cas A : Colonne 0 = ID numérique (ex: 14)
            if c0.isdigit():
                student = students_by_id.get(int(c0))
            # Cas B : Colonne 0 = Email
            elif '@' in c0:
                student = students_by_email.get(c0.lower())
            # Cas C : Colonne 0 = Nom, Colonne 1 = Prénom
            elif len(row) >= 3 and row[2] is not None:
                student = students_by_name.get(f"{c0.lower()} {c1.lower()}") or students_by_name_rev.get(f"{c0.lower()} {c1.lower()}")

            # Trouver la note : dernière colonne numérique
            if score_val is None:
                for val in reversed(row):
                    if val is not None and str(val).strip() != "":
                        try:
                            _s = float(str(val).replace(',', '.'))
                            if 0 <= _s <= 20:
                                score_val = _s
                                break
                        except ValueError:
                            pass

        if not student:
            errors += 1
            info_label = f"Ligne {row_idx}: Étudiant non identifié ({row[:3]})"
            error_details.append(info_label)
            continue

        # Convertir et vérifier la note
        try:
            if score_val is None:
                raise ValueError("Note absente")
            score = float(str(score_val).replace(',', '.'))
            if score < 0 or score > 20:
                raise ValueError("Note hors plage (0-20)")
        except (TypeError, ValueError) as err_score:
            errors += 1
            error_details.append(f"Ligne {row_idx} ({student.first_name} {student.last_name}): note invalide '{score_val}'")
            continue

        # Upsert
        existing = db.query(Grade).filter(
            Grade.student_id == student.id,
            Grade.module_id == module_id,
            Grade.grade_type == gtype,
            Grade.academic_year == academic_year,
        ).first()

        if existing:
            existing.score = score
            existing.uploaded_by_id = current_user.id
            updated += 1
        else:
            new_grade = Grade(
                student_id=student.id,
                module_id=module_id,
                grade_type=gtype,
                score=score,
                academic_year=academic_year,
                uploaded_by_id=current_user.id,
            )
            db.add(new_grade)
            imported += 1

    db.commit()

    total = imported + updated
    return GradeBulkUploadResponse(
        imported=imported,
        updated=updated,
        errors=errors,
        error_details=error_details[:20],
        message=f"{total} note(s) traitée(s) avec succès ({imported} créées, {updated} modifiées, {errors} non résolues)."
    )


@router.get("/me/grades", response_model=list[GradeRead])
def get_my_uploaded_grades(
    module_id: Optional[int] = None,
    grade_type: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(_require_teacher),
):
    """Retourne toutes les notes uploadées par le professeur connecté."""
    from sqlalchemy.orm import joinedload
    query = db.query(Grade).filter(Grade.uploaded_by_id == current_user.id)
    if module_id:
        query = query.filter(Grade.module_id == module_id)
    if grade_type:
        try:
            gtype = GradeType(grade_type.lower())
            query = query.filter(Grade.grade_type == gtype)
        except ValueError:
            pass

    grades = query.options(
        joinedload(Grade.student),
        joinedload(Grade.module),
        joinedload(Grade.uploaded_by),
    ).order_by(Grade.module_id, Grade.grade_type, Grade.student_id).all()
    return grades
